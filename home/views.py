
from django.shortcuts import render, redirect
from django import forms
from django.http import JsonResponse
from django.contrib.auth.decorators import user_passes_test
from django.urls import reverse
from django.contrib import messages
from .decorators import check_client
from django.contrib.auth.decorators import permission_required
from django.db.models import Q
from .models import Cliente, Offerta, ModelloPresentazioneOfferta, Variazione, DocumentoCompilato

from django.http import HttpResponse
from django.template.loader import render_to_string
import pdfkit

PDF_OPTIONS = {
    'page-size': 'A4',
    'margin-top': '0.5in',
    'margin-right': '0.5in',
    'margin-bottom': '0.5in',
    'margin-left': '0.5in',
}

def index(request):
    if 'client_id' in request.session:
        return redirect(reverse('home:offerte'))
    if request.method == 'POST' and 'send' in request.POST:
        email = request.POST['email']
        clienti = Cliente.objects.filter(email=email)
        for cliente in clienti:
            cliente.send_token(request) # passo la request per avere l'host
        messages.success(request, 'Ti abbiamo inviato un’email con il link per accedere al tuo account.')
        return redirect(request.path)
    return render(request, 'home/index.html')

@check_client
def offerta_view(request, pk, pdf=False, check_session=True):
    from django.template.defaultfilters import slugify
    offerta = Offerta.objects.get(pk=pk)

    if request.user.is_authenticated:
        cliente = offerta.cliente
        check_session = False
    else:
        cliente = Cliente.objects.get(pk=request.session['client_id'])
    if cliente == offerta.cliente or not check_session:
        if pdf:
            from PyPDF2 import PdfWriter, PdfReader
            import io

            out = None

            # Aggiungo la presentazione se c'è
            if offerta.presentazione:
                html_presentazione = render_to_string('home/offerta_presentazione.html', {'object': offerta})
                pdf_presentazione = pdfkit.from_string(html_presentazione, False, options=PDF_OPTIONS)
                out = PdfWriter()
                for page in PdfReader(io.BytesIO(pdf_presentazione)).pages:
                    out.add_page(page)

            html = render_to_string('home/offerta.html', {'object': offerta, 'pdf': pdf, 'info_utente': load_info_utente(), 'request': request})
            if 'html' in request.GET:
                return HttpResponse(html)
            
            pdf = pdfkit.from_string(html, False, options=PDF_OPTIONS)



            response = HttpResponse(content_type='application/pdf')
            if out:
                for page in PdfReader(io.BytesIO(pdf)).pages:
                    out.add_page(page)

                out.write(response)
            else:
                response.write(pdf)

            filename = (slugify(offerta.titolo) or 'offerta') + '.pdf'
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        else:
            return render(request, 'home/offerta.html', {'object': offerta})
    return redirect('/')

@check_client
def offerta_presentazione(request, pk):
    cliente = Cliente.objects.get(pk=request.session['client_id'])
    offerta = Offerta.objects.get(pk=pk)

    if cliente == offerta.cliente:
        html = render_to_string('home/offerta_presentazione.html', {'object': offerta})

        pdf = pdfkit.from_string(html, False, options=PDF_OPTIONS)
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="presentazione-offerta.pdf"'
        return response

    return redirect('/')


@check_client
def accettazione_offerta(request, pk, compilato_id):
    from PyPDF2 import PdfWriter, PdfReader
    import io
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from operator import attrgetter
    from django.template.defaultfilters import slugify
    


    cliente = Cliente.objects.get(pk=request.session['client_id'])
    offerta = Offerta.objects.get(pk=pk)
    allegato = DocumentoCompilato.objects.get(pk=compilato_id)
    if cliente == offerta.cliente:
        response = HttpResponse(content_type='application/pdf')
        filename = slugify(allegato) + '.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        if allegato.campi.exists():
            existing_pdf = PdfReader(open(allegato.documento.file.path, "rb"))
            output = PdfWriter()

            for i, page in enumerate(existing_pdf.pages):

                compilazione = io.BytesIO()
                campi_pagina = allegato.campi.filter(n_pagina=i+1)
                if campi_pagina.count() == 0:
                    output.add_page(page)
                    continue

                can = canvas.Canvas(compilazione, pagesize=A4)

                for campo in campi_pagina:
                    can.setFont("Helvetica", 9)
                    can.drawString(campo.posizione_x*cm, campo.posizione_y*cm, attrgetter(campo.campo)(offerta))
                can.save()

                #move to the beginning of the StringIO buffer
                compilazione.seek(0)

                # create a new PDF with Reportlab
                new_pdf = PdfReader(compilazione)
                # add the "watermark" (which is the new pdf) on the existing page
                page.merge_page(new_pdf.pages[0])
                output.add_page(page)

            output.write(response)

            return response
    
        else:
            response.write(allegato.documento.file.read())
            return response

    return redirect('/')

@permission_required('change_offerta')
def offerta_presentazione_edit(request, pk, pdf=False):
    offerta = Offerta.objects.get(pk=pk)
    modelli = ModelloPresentazioneOfferta.objects.all()

    if request.method == 'POST' and 'save' in request.POST:
        modello = ModelloPresentazioneOfferta.objects.get(pk=request.POST['modello'])
        offerta.presentazione = modello.testo
        offerta.save()
        return redirect(request.path)
    
    return render(request, 'home/offerta_presentazione_edit.html', {'object': offerta, 'modelli': modelli})


@check_client
def offerte_listing(request):
    cliente = Cliente.objects.get(pk=request.session['client_id'])
    if request.method == 'POST':
        if 'visibile' in request.POST:
            offerta = Offerta.objects.get(cliente=cliente, pk=request.POST['visibile'])
            offerta.visibile = not offerta.visibile
            offerta.save()
            return redirect(request.path)

    offerte = cliente.offerte.all()
    context = {
        'cliente': cliente, # per il menu
        'offerte': offerte
    }
    return render(request, 'home/offerte.html', context)

@check_client
def dati_cliente(request):
    cliente = Cliente.objects.get(pk=request.session['client_id'])
    if request.method == 'POST':
        if 'save' in request.POST:
            Variazione.objects.create(cliente=cliente, descrizione=request.POST['descrizione'])
    context = {
        'cliente': cliente
    }
    return render(request, 'home/dati_cliente.html', context)


@check_client
def documenti_cliente(request):
    from django.core.files.base import File
    from wagtail.documents.models import Document
    from wagtail.models import Collection
    import os
    from .models import DocumentoCliente

    cliente = Cliente.objects.get(pk=request.session['client_id'])
 #   docs_da_firmare = AccettazioneOfferta.objects.filter(offerta__cliente=cliente)

    if request.method == 'POST' and 'upload' in request.POST:
        for f in request.FILES.getlist('files'):
            filename = f.name
            doc_file = File(f, name=os.path.basename(filename))
            doc = Document(
                title=filename,
                file=doc_file,
            )
            doc.save()
            DocumentoCliente.objects.create(cliente=cliente, document=doc)
        return redirect(reverse('home:documenti-cliente'))
    
    if request.method == 'POST' and 'delete' in request.POST:
        DocumentoCliente.objects.filter(pk=request.POST['delete'], cliente=cliente, privato=False, document__uploaded_by_user=None).delete()
        return redirect(reverse('home:documenti-cliente'))
    
    # if request.method == 'POST' and 'signed' in request.POST:
    #     collection = Collection.objects.filter(name='Documenti firmati').first()
    #     allegato = AccettazioneOfferta.objects.filter(pk=request.POST['signed'], offerta__cliente=cliente).first()
    #     if allegato:
    #         filename = request.FILES['file'].name
    #         doc_file = File(request.FILES['file'], name=os.path.basename(filename))
    #         doc = Document(
    #             title=allegato.documento.documento.title + ' firmato',
    #             file=doc_file,
    #             collection=collection,
    #         )
    #         doc.save()
    #         allegato.firmato = doc
    #         allegato.save()
    #     return redirect(reverse('home:documenti-cliente'))




    documenti = cliente.documenti.filter(privato=False)
    context = {
        'cliente': cliente,
        'documenti': documenti,
    #    'docs_da_firmare':docs_da_firmare,
    }
    return render(request, 'home/documenti_cliente.html', context)


@check_client
def documenti_prodotti(request):
    cliente = Cliente.objects.get(pk=request.session['client_id'])
 
    documenti = cliente.documenti_prodotto.all()
    print(documenti)
    print(cliente)
    context = {
        'cliente': cliente,
        'documenti': documenti
    }
    return render(request, 'home/documenti_prodotti.html', context)



def cliente_start(request, pk, token, section=None, offerta_id=None):
    from .models import Cliente
    cliente = Cliente.objects.get(pk=pk)
    if token == cliente.token:
        # Link diretto all'offerta in pdf, non metto in sessione
        if offerta_id:
            # Verifico che l'offerta sia del cliente
            offerta = Offerta.objects.get(cliente=cliente, pk=offerta_id)
            return offerta_view(request, offerta.pk, pdf=True, check_session=False)
        
        request.session['client_id'] = cliente.pk
        if section:
            return redirect('/' + section)
        return redirect(reverse('home:offerte'))
    return redirect('/')


def prodotti_search(request):
    from .models import Prodotto
    term = request.GET.get('q', '')
    prodotti = Prodotto.objects.filter(Q(codice__icontains=term) | Q(descrizione__icontains=term), attivo=True)[:20]
    results = [prodotto.to_json() for prodotto in prodotti]
    return JsonResponse(results, safe=False)


def comuni_search(request):
    from .models import Comune
    term = request.GET.get('q', '')
    if term:
        comuni = Comune.objects.filter(nome__icontains=term)[:20]
        results = [comune.to_json() for comune in comuni]
    else:
        results = []
    return JsonResponse(results, safe=False)


class ImportaClienteForm(forms.Form):
    url = forms.URLField(initial='https://aziende.clientiperte.com/DL?r=TB11FFLMMR1L')



def admin_importa_cliente(request):
    from pyquery import PyQuery as pq
    from .models import Cliente, Segnalazione, Gruppo
    from django.contrib import messages

    form = ImportaClienteForm(request.POST or None)

    if request.method == 'POST':
        if form.is_valid():
            url = form.cleaned_data['url']
            print(url)
            d = pq(url=url)
            section_titles = d('.text-info')

            data = {}
            for section_title in section_titles:
                section_value = d(section_title).closest('div').remove('.text-info').text().strip()
                data[section_title.text.lower()] = section_value

            print(data)
            segnalazione = Segnalazione.objects.filter(nome='ClientiPerTe').first()
            gruppo = Gruppo.objects.filter(nome__iexact='Nuovi contatti').first()

            cliente = Cliente(ragsoc=data.get('cognome', '').title() + ' ' + data.get('nome', '').title(), email=data.get('email', ''), telefono=data.get('telefono', ''), citta=data.get('comune', ''), provincia=data.get('provincia', '').upper()[:2], note=data.get('testo della richiesta', ''), segnalazione=segnalazione, gruppo=gruppo)

            cliente.save()
            messages.success(request, 'Cliente importato con successo')
            cliente_edit_url = f'/admin/snippets/home/cliente/edit/{cliente.pk}/'
            return redirect(cliente_edit_url)
    return render(request, 'modeladmin/importa_cliente.html', {
        'form': form
    })


def add_listino_1(request):
    from .forms import ListinoForm
    context = {}
    form = ListinoForm(request.POST or None, request.FILES or None)
    if request.method == 'POST':
        if form.is_valid():
            listino = form.save()
            return redirect(reverse('admin_add_listino_2', args=[listino.pk]))
        
    context['form'] = form
    context['title'] = 'Aggiungi listino'

    return render(request, 'modeladmin/add_listino.html', context=context)



def add_listino_2(request, id):
    """Permette di selezionare le colonne da importare"""
    import openpyxl
    from .models import Listino, Prodotto

    listino = Listino.objects.get(pk=id)

    wb = openpyxl.load_workbook(listino.file.path, data_only=True)
    sh = wb.active
    table_data = list(sh.iter_rows(min_row=1, values_only=True))
   
    if request.method == 'POST':
        if 'select' in request.POST:
            codice_index, descrizione_index, prezzo_index = -1, -1, -1
            for k, v in request.POST.items():
                if k.startswith('col-'):
                    index = int(k.split('-')[-1])
                    if v == 'codice':
                        codice_index = index
                    if v == 'descrizione':
                        descrizione_index = index
                    if v == 'prezzo':
                        prezzo_index = index


            if codice_index == -1 or descrizione_index == -1 or prezzo_index == -1:
                messages.error(request, 'Devi selezionare tutte e tre le colonne (Codice, Descrizione e Prezzo)')
            else:
                context = {'codice_index': codice_index, 'descrizione_index': descrizione_index, 'prezzo_index': prezzo_index, 'listino': listino}
                prodotti = []
                for row in table_data:
                    if row[codice_index] and row[descrizione_index] and type(row[prezzo_index]) in (float, int):
                        prodotto = {'codice': row[codice_index], 'descrizione': row[descrizione_index], 'prezzo': row[prezzo_index]}
                        prodotti.append(prodotto)
                context['prodotti'] = prodotti
                return render(request, 'modeladmin/add_listino_verify.html', context) 

        if 'save' in request.POST:
            codice_index = int(request.POST['codice_index'])
            descrizione_index = int(request.POST['descrizione_index'])
            prezzo_index = int(request.POST['prezzo_index'])
            context = {'listino': listino}
            prodotti = 0
            for row in table_data:
                if row[codice_index] and row[descrizione_index] and type(row[prezzo_index]) in (float, int):
                    prodotto = Prodotto.objects.create(listino=listino, codice=row[codice_index], descrizione=row[descrizione_index], prezzo=row[prezzo_index])
                    prodotti += 1
            context['prodotti'] = prodotti
            return render(request, 'modeladmin/add_listino_done.html', context=context)


    table_data = [row for row in table_data if any(row)][:20]
    context = {'table_data': table_data, 'listino': listino}
    return render(request, 'modeladmin/add_listino_columns.html', context=context)



 
def load_info_utente():
    from .models import InformazioneUtente
    info_utente = {}
    records = InformazioneUtente.objects.all()

    for r in records:
        if r.tipo == 'C':
            value = r.valore_char
        if r.tipo == 'T':
            value = r.valore_text
        if r.tipo == 'I':
            value = r.valore_image
        if r.tipo == 'B':
            value = r.valore_boolean
        info_utente[r.codice] = value
        
    return info_utente



# Link da mandare l'installatore che mostra tutti i dati di un'offerta
def installatore_start(request, pk, token):
    from .models import Offerta
    offerta = Offerta.objects.get(pk=pk)
    if token == offerta.token:
        documenti = offerta.cliente.documenti.filter(installatore=True)
        context = {
            'offerta': offerta,
            'documenti': documenti,
            'installatore': True
        }
        return render(request, 'home/installatore.html', context)
    return redirect('/')